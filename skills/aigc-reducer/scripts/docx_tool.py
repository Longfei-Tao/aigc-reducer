#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""docx_tool.py — 论文降 AIGC 率工具

  split   docx -> 段落清单 + 章节树 + 分批方案 (JSON)
  batch   从 split 结果里取一批段落，拼成可直接发送的改写请求
  apply   模型回复 -> edits.json（段落数不符直接拒绝）
  status  查看各批完成情况，中断后续跑
  write   改写结果 -> 新 docx（对照表仅按需导出）
  check   改写前后硬指标自检

依赖：python-docx
    安装：python -m pip install python-docx

示例：
  python docx_tool.py split 论文.docx -o work/blocks.json --batch-chars 1200
  python docx_tool.py batch work/blocks.json --id 3 --prompt references/prompt.md -o work/batch3.txt
  python docx_tool.py apply work/blocks.json --batch 3 --input work/reply3.txt -o work/edits.json
  python docx_tool.py status work/blocks.json
  python docx_tool.py check 论文.docx work/edits.json
  python docx_tool.py write 论文.docx work/edits.json -o 论文_降AI版.docx
"""

import argparse
import hashlib
import json
import os
import re
import sys
import tempfile
from collections import Counter
from datetime import datetime

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    sys.exit(
        "缺少依赖 python-docx。\n"
        "请执行：python -m pip install python-docx"
    )

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

# ---------------------------------------------------------------- 段落分类

REF_MARKERS = ("参考文献", "references", "致谢", "acknowledg", "附录", "appendix")
# 手打编号标题：第X章 / 一、 / （一） / 1.1 / 1、
MANUAL_HEADING = re.compile(
    r"^\s*(?:"
    r"第[一二三四五六七八九十百]+[章节篇部]"
    r"|[一二三四五六七八九十]+[、\.]"
    r"|[（(][一二三四五六七八九十]+[）)]"
    r"|\d+(?:\.\d+)*[\.、\s]"
    r")"
)
# 以年份开头的正文（「2023 年共收集……」）不是标题，避免把正文判成标题而漏改
YEAR_START = re.compile(r"^(?:19|20)\d{2}\s*年")
CAPTION = re.compile(r"^\s*(图|表|Fig|Table)\s*[\d\-—\.]")
BANNED_SUBJECTIVE = ("我觉得", "我认为", "在我看来", "窃以为", "笔者觉得")
AI_CLICHE = ("综上所述", "值得注意的是", "不难发现", "由此可见", "显著提升", "至关重要")
# 模型把指令回显进正文的典型特征。出现任意一条即为明确故障，不是措辞问题
ECHO_MARKERS = ("修改段落", "在不大幅改动原文总字数的情况下")

NUM_RE = re.compile(r"\d+(?:\.\d+)?")
VALUE_RE = re.compile(r"[<>≤≥=]?\s*[+−-]?\d+(?:\.\d+)?(?:[eE][+−-]?\d+)?\s*[%％]?")
# 只认方括号引用（GB/T 7714 的 [1]、[1-3]、[1,2]）。
# 中文论文里「（2021）」「（1）」多为年份括注或列表编号，误当引用会产生假的「引用丢失」fail
CITE_RE = re.compile(r"[\[［【]\s*\d+(?:\s*[-,，–—~]\s*\d+)*\s*[\]］】]")

# 只有 para 允许被改写；其余类型一律拦截，避免标题、图注被正文覆盖
EDITABLE_KINDS = ("para",)
KIND_LABEL = {
    "para": "正文",
    "heading": "标题",
    "caption": "图表题注",
    "ref": "参考文献/致谢",
    "formula": "公式",
    "skip": "空白或图片",
    "complex": "图片/域/修订等复杂对象",
}


def has_complex_content(p):
    # r.text 赋值会清除 run 的其他 XML 子节点；这类结构不可进入替换路径。
    tags = ("w:drawing", "w:pict", "w:object", "w:fldChar", "w:instrText",
            "w:fldSimple", "w:sdt", "w:ins", "w:del", "w:commentRangeStart",
            "w:commentRangeEnd", "w:commentReference", "w:txbxContent")
    return any(run_has_note_ref(r) and r.text for r in p.runs) or any(
        any(True for _ in p._p.iter(qn(tag))) for tag in tags
    )


def has_image(p):
    xml = p._element.xml
    return "<w:drawing" in xml or "<w:pict" in xml or "<pic:pic" in xml


def has_formula(p):
    xml = p._element.xml
    return "<m:oMath" in xml or "<m:oMathPara" in xml


def style_name(p):
    try:
        return p.style.name or ""
    except Exception:
        return ""


def is_heading_style(p):
    name = style_name(p)
    if not name:
        return False
    n = name.lower()
    return n.startswith("heading") or n.startswith("title") or "标题" in name


def looks_like_manual_heading(text):
    t = text.strip()
    if not t or len(t) > 50:
        return False
    if t.endswith(("。", "！", "？", ".", ";", "；")):
        return False
    if YEAR_START.match(t):
        return False
    return bool(MANUAL_HEADING.match(t))


def classify(p, in_refs):
    """返回 (kind, text)"""
    text = p.text or ""
    stripped = text.strip()

    if has_image(p) and not stripped:
        return "skip", text
    if has_formula(p):
        return "formula", text
    if has_complex_content(p):
        return "complex", text
    if not stripped:
        return "skip", text

    if is_heading_style(p):
        return "heading", text
    if CAPTION.match(stripped):
        return "caption", text
    if looks_like_manual_heading(stripped):
        return "heading", text
    if in_refs:
        return "ref", text
    return "para", text


def scan_blocks(doc):
    blocks = []
    in_refs = False
    for i, p in enumerate(doc.paragraphs):
        text = p.text or ""
        low = text.strip().lower()
        marker_heading = low in ("参考文献", "references", "致谢", "acknowledgments",
                                  "acknowledgements", "附录", "appendix")
        if not in_refs and (marker_heading or (is_heading_style(p) and any(m in low for m in REF_MARKERS))):
            in_refs = True
        kind, txt = classify(p, in_refs)
        if marker_heading:
            kind = "heading"
        blocks.append(
            {
                "id": i,
                "kind": kind,
                "style": style_name(p),
                "chars": len(txt.strip()),
                "runs": len(p.runs),
                "text": txt,
            }
        )
    return blocks


def build_sections(blocks):
    """章节树：每个 heading 后面到下一个 heading 之前的正文统计"""
    sections = []
    cur = None
    for b in blocks:
        if b["kind"] == "heading":
            cur = {
                "heading_id": b["id"],
                "heading": b["text"].strip(),
                "para_ids": [],
                "chars": 0,
            }
            sections.append(cur)
        elif b["kind"] == "para" and cur is not None:
            cur["para_ids"].append(b["id"])
            cur["chars"] += b["chars"]
    return sections


def build_batches(blocks, limit):
    batches = []
    cur_ids, cur_chars, cur_sec = [], 0, "（文首）"

    def flush():
        nonlocal cur_ids, cur_chars
        if cur_ids:
            batches.append(
                {
                    "batch_id": len(batches) + 1,
                    "section": cur_sec,
                    "para_ids": list(cur_ids),
                    "chars": cur_chars,
                }
            )
            cur_ids, cur_chars = [], 0

    for b in blocks:
        if b["kind"] == "heading":
            flush()
            cur_sec = b["text"].strip()[:60]
            continue
        if b["kind"] != "para":
            flush()
            continue
        if cur_ids and cur_chars + b["chars"] > limit:
            flush()
        cur_ids.append(b["id"])
        cur_chars += b["chars"]
    flush()
    return batches


# ---------------------------------------------------------------- split

def cmd_split(args):
    if args.batch_chars <= 0:
        sys.exit("[错误] --batch-chars 必须是正整数")
    doc = open_docx(args.docx)
    blocks = scan_blocks(doc)
    sections = build_sections(blocks)
    batches = build_batches(blocks, args.batch_chars)

    eligible = [b for b in blocks if b["kind"] == "para"]
    tables = len(doc.tables)
    table_chars = 0
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                table_chars += len(cell.text or "")

    out = {
        "source": os.path.abspath(args.docx),
        "source_sha256": file_digest(args.docx),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "batch_chars": args.batch_chars,
        "summary": {
            "para_total": len(blocks),
            "eligible_paras": len(eligible),
            "eligible_chars": sum(b["chars"] for b in eligible),
            "heading_count": sum(1 for b in blocks if b["kind"] == "heading"),
            "ref_paras": sum(1 for b in blocks if b["kind"] == "ref"),
            "formula_paras": sum(1 for b in blocks if b["kind"] == "formula"),
            "caption_paras": sum(1 for b in blocks if b["kind"] == "caption"),
            "tables": tables,
            "table_chars": table_chars,
            "batches": len(batches),
        },
        "sections": sections,
        "batches": batches,
        "blocks": blocks,
    }

    write_json(args.output, out)
    s = out["summary"]
    print(f"[split] {os.path.basename(args.docx)}")
    print(f"  段落总数      {s['para_total']}")
    print(f"  可改写正文     {s['eligible_paras']} 段 / {s['eligible_chars']} 字")
    print(f"  标题 / 参考文献 / 公式 / 题注  {s['heading_count']} / {s['ref_paras']} / {s['formula_paras']} / {s['caption_paras']}")
    print(f"  表格          {s['tables']} 个（{s['table_chars']} 字，不参与改写）")
    print(f"  分批          {s['batches']} 批（每批 ≤{args.batch_chars} 字）")
    print(f"  输出          {args.output}")


# ---------------------------------------------------------------- batch

def read_prompt(path):
    """取第一个 --- 分隔线之后的全部内容（含「修改段落:」）"""
    with open(path, "r", encoding="utf-8") as f:
        raw = f.read()
    m = re.search(r"(?m)^---\s*$", raw)
    if not m:
        sys.exit(f"[错误] {path} 中没有找到 --- 分隔线，无法提取改写指令")
    body = raw[m.end():].lstrip("\n")
    if "修改段落" not in body:
        print(
            f"[警告] 从 {os.path.basename(path)} 提取的指令里没有「修改段落:」，"
            "请确认 --prompt 指向 references/prompt.md",
            file=sys.stderr,
        )
    return body


def match_section(sections, keyword):
    kw = keyword.strip()
    hits = [s for s in sections if s["heading"].strip() == kw]
    if not hits:
        hits = [s for s in sections if kw and (kw in s["heading"] or s["heading"] in kw)]
    if not hits:
        # 去编号后再试
        norm = re.sub(r"^[\d\.\s、]+", "", kw)
        if norm:
            hits = [s for s in sections if norm in s["heading"]]
    return hits


def cmd_batch(args):
    data = load_json(args.blocks)
    block_map = {b["id"]: b for b in data["blocks"]}
    batches = data["batches"]

    if args.id is not None:
        hit = [b for b in batches if b["batch_id"] == args.id]
        if not hit:
            sys.exit(f"[错误] 没有第 {args.id} 批，共 {len(batches)} 批")
        ids = hit[0]["para_ids"]
        section = hit[0]["section"]
        label = f"批次 {args.id} / 共 {len(batches)} 批"
    elif args.ids:
        raw = args.ids.replace("，", ",")
        try:
            ids = [int(x) for x in raw.split(",") if x.strip()]
        except ValueError:
            sys.exit(f"[错误] --ids 只接受逗号分隔的整数，收到：{args.ids}")
        section = ""
        label = "自定义段落"
    elif args.section:
        hits = match_section(data["sections"], args.section)
        if not hits:
            names = [s["heading"] for s in data["sections"] if s["heading"]]
            sys.exit(
                f"[错误] 没有匹配到章节「{args.section}」。\n"
                f"可用章节：\n  " + "\n  ".join(names[:40])
            )
        if len(hits) > 1:
            names = [s["heading"] for s in hits]
            sys.exit(f"[错误] 「{args.section}」匹配到多节：{names}，请用 --id 或 --ids 明确范围")
        ids = list(hits[0]["para_ids"])
        section = hits[0]["heading"]
        label = f"章节 {section}"
    else:
        sys.exit("[错误] 需要 --id / --ids / --section 之一，三者选一")

    missing = [i for i in ids if i not in block_map]
    if missing:
        sys.exit(f"[错误] blocks 里不存在这些段落 id：{missing}")
    blocked = [i for i in ids if block_map[i]["kind"] not in EDITABLE_KINDS]
    if blocked:
        detail = "\n  ".join(
            f"{i} -> {KIND_LABEL.get(block_map[i]['kind'], block_map[i]['kind'])}：{block_map[i]['text'][:30]}"
            for i in blocked
        )
        sys.exit(f"[错误] 这些 id 不是正文段落，不能参与改写：\n  {detail}")
    if not ids:
        sys.exit("[错误] 这一批没有可改写的正文段落")
    if len(ids) != len(set(ids)):
        sys.exit("[错误] 段落 id 重复，拒绝处理")

    texts = [block_map[i]["text"].strip() for i in ids]
    chars = sum(len(t) for t in texts)
    body = "\n\n".join(texts)

    if args.prompt:
        payload = read_prompt(args.prompt).rstrip() + "\n" + body + "\n"
    else:
        payload = body + "\n"

    if args.output:
        os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(payload)

    print(f"[batch] {label}", file=sys.stderr)
    print(f"  段落   {len(ids)} 段 / {chars} 字", file=sys.stderr)
    print(f"  id 顺序 {','.join(str(i) for i in ids)}", file=sys.stderr)
    if args.output:
        print(f"  请求文本 -> {args.output}（可直接复制发送）", file=sys.stderr)
    else:
        print("  提示：加 -o 保存为纯请求文本，避免和上面的提示混在一起", file=sys.stderr)

    if not args.output:
        sys.stdout.write(payload)


# ---------------------------------------------------------------- write

def load_edits(path):
    data = load_json(path)
    if isinstance(data, dict) and "schema_version" in data:
        if data["schema_version"] != 1 or not isinstance(data.get("edits"), dict):
            sys.exit("[错误] 不支持的 edits 版本或结构")
        data = data["edits"]
    if isinstance(data, dict):
        items = data
    elif isinstance(data, list):
        items = {}
        for it in data:
            if not isinstance(it, dict) or "id" not in it or "text" not in it:
                sys.exit('[错误] edits 列表模式下每项必须是 {"id": 段落号, "text": "新文本"}')
            items[str(it["id"])] = it["text"]
    else:
        sys.exit("[错误] edits 文件格式应为 {id: 新文本} 或 [{id, text}]")

    bad_keys = [k for k in items if not str(k).strip().lstrip("-").isdigit()]
    if bad_keys:
        sys.exit(
            "[错误] edits 的 key 必须是段落整数 id，这些不是："
            f"{bad_keys[:10]}\n"
            "段落 id 请从 blocks.json 的 blocks[].id 取，或先用 batch 子命令确认。"
        )
    bad_vals = [k for k, v in items.items() if not isinstance(v, str) or not v.strip()]
    if bad_vals:
        sys.exit(f"[错误] 这些段落的改写文本为空或不是字符串：{bad_vals[:10]}")
    return {int(k): v for k, v in items.items()}


def same_file(a, b):
    try:
        return os.path.samefile(a, b)
    except OSError:
        return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def check_targets(paras, edits, block_map, allow_non_para):
    """拦截非正文段落，避免把标题/图注/参考文献覆盖成正文"""
    out_of_range, blocked = [], []
    for pid in sorted(edits):
        if pid < 0 or pid >= len(paras):
            out_of_range.append(pid)
            continue
        kind = block_map.get(pid, {}).get("kind", "para")
        if kind not in EDITABLE_KINDS and not allow_non_para:
            blocked.append((pid, kind, block_map.get(pid, {}).get("text", "")))
    return out_of_range, blocked


def run_has_note_ref(r):
    """脚注/尾注引用所在的 run —— 改写时必须原样保留，否则脚注会消失"""
    xml = r._element.xml
    return "footnoteReference" in xml or "endnoteReference" in xml


def format_sig(r):
    """run 的格式签名，用于判断段落内是否存在多种格式"""
    f = r.font
    try:
        color = str(f.color.rgb) if f.color and f.color.rgb else None
    except Exception:
        color = None
    return (bool(r.bold), bool(r.italic), bool(r.underline), color, f.size)


def cmd_write(args):
    verify_blocks_source(args.blocks, args.docx)
    verify_edits_source(args.edits, args.docx)
    if os.path.exists(args.output):
        sys.exit("[错误] 输出文件已存在，请使用新的 -o 文件名")
    if same_file(args.docx, args.output):
        sys.exit(
            "[错误] 输出文件与源文件相同，会覆盖原稿。\n"
            "本工具不覆盖原文件，请换一个 -o 文件名，例如 论文_降AI版.docx"
        )

    doc = open_docx(args.docx)
    blocks = scan_blocks(doc)
    block_map = {b["id"]: b for b in blocks}
    edits = load_edits(args.edits)
    paras = doc.paragraphs

    out_of_range, blocked = check_targets(paras, edits, block_map, args.allow_non_para)
    complex_ids = [pid for pid in edits if 0 <= pid < len(paras) and has_complex_content(paras[pid])]
    if complex_ids:
        sys.exit(f"[错误] 段落含不可安全替换的复杂对象，拒绝回写：{complex_ids}")
    for pid in out_of_range:
        print(f"  警告：段落 id {pid} 越界，已跳过", file=sys.stderr)
    if blocked:
        detail = "\n  ".join(
            f"{i} -> {KIND_LABEL.get(k, k)}：{t[:30]}" for i, k, t in blocked
        )
        sys.exit(
            f"[错误] {len(blocked)} 个 id 指向的不是正文段落，已全部拒绝回写：\n  {detail}\n"
            "这类段落（标题、图注、参考文献、公式）被正文覆盖后很难复原。\n"
            "请核对 blocks.json 的 blocks[].id；确需改写请加 --allow-non-para。"
        )

    echo_ids = [(pid, echo_hits(t)) for pid, t in sorted(edits.items()) if echo_hits(t)]
    if echo_ids:
        detail = "\n  ".join(f"{pid} -> 含「{','.join(h)}」：{edits[pid][:40]}" for pid, h in echo_ids)
        sys.exit(
            f"[错误] {len(echo_ids)} 段把改写指令回显进了正文，拒绝回写：\n  {detail}\n"
            "说明上一批没有得到「只包含段落正文」的回复。请重新生成该批，"
            "或检查 apply 是否把整段回复当成了一个段落。"
        )

    applied, missing, no_run = 0, [], 0
    skipped_link, multi_fmt, note_kept, note_head = [], [], [], []
    written = set()

    for pid, new_text in sorted(edits.items()):
        if pid < 0 or pid >= len(paras):
            missing.append(pid)
            continue
        p = paras[pid]

        if "<w:hyperlink" in p._element.xml and not args.force:
            skipped_link.append(pid)
            continue

        runs = p.runs
        if not runs:
            p.add_run(new_text)
            no_run += 1
            applied += 1
            written.add(pid)
            continue

        if len({format_sig(r) for r in runs}) > 1:
            multi_fmt.append(pid)

        editable = [r for r in runs if not run_has_note_ref(r)]
        if not editable:
            p.add_run(new_text)
            no_run += 1
            applied += 1
            written.add(pid)
            continue

        if run_has_note_ref(runs[0]):
            note_head.append(pid)
        if len(editable) != len(runs):
            note_kept.append(pid)

        editable[0].text = new_text
        for r in editable[1:]:
            r.text = ""
        applied += 1
        written.add(pid)

    if edits and not written:
        detail = []
        if missing:
            detail.append(f"{len(missing)} 个段落 id 越界（此文档共 {len(paras)} 段，id 从 0 开始计）")
        if skipped_link:
            detail.append(f"{len(skipped_link)} 段含超链接被跳过")
        sys.exit(
            "[错误] 没有任何段落被回写，未生成新文件。\n  "
            + "\n  ".join(detail)
            + "\n  请核对 edits.json 的 key 是否取自 blocks.json 的 blocks[].id。"
        )
    if not edits:
        print("[提示] edits 里没有任何段落，输出的是原文档副本")

    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or ".", exist_ok=True)
    doc.save(args.output)

    print(f"[write] 已回写 {applied} 段 -> {args.output}")
    if no_run:
        print(f"  提示：{no_run} 段原本无可写 run，已新建（继承段落样式）")
    if missing:
        print(f"  警告：{len(missing)} 个段落 id 越界，已跳过：{missing[:10]}")
    if skipped_link:
        print(f"  [跳过] {len(skipped_link)} 段含超链接，未改写（强行改写会让链接文字重复）。需手工处理：{skipped_link[:10]}")
    if multi_fmt:
        print(f"  [提示] {len(multi_fmt)} 段内含多种字符格式（如局部加粗），改写后统一为段首格式，局部强调会丢失：{multi_fmt[:10]}")
    if note_kept:
        print(f"  [提示] {len(note_kept)} 段含脚注/尾注引用，引用已保留，但位置可能随语句调整而偏移：{note_kept[:10]}")
    if note_head:
        print(f"  [警告] {len(note_head)} 段脚注引用位于段首，改写文字可能排到引用之后，建议手工核对：{note_head[:10]}")

    if args.diff_md:
        real_edits = {k: v for k, v in edits.items() if k in written}
        write_diff_md(args.docx, real_edits, args.diff_md)
        print(f"  对照表 -> {args.diff_md}（{len(real_edits)} 段）")


def write_diff_md(src, edits, path):
    doc = open_docx(src)
    paras = doc.paragraphs
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    lines = [
        "# 改动对照表",
        "",
        f"- 源文件：{os.path.basename(src)}",
        f"- 生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"- 改动段落：{len(edits)} 段",
        "",
        "| 段落 | 原文字数 | 新文字数 |",
        "|---|---|---|",
    ]
    for pid in sorted(edits):
        if 0 <= pid < len(paras):
            old = (paras[pid].text or "").strip()
            new = edits[pid].strip()
            lines.append(f"| {pid} | {len(old)} | {len(new)} |")
    lines.append("")
    for pid in sorted(edits):
        if not (0 <= pid < len(paras)):
            continue
        old = (paras[pid].text or "").strip()
        new = edits[pid].strip()
        lines += [
            f"## 段落 {pid}",
            "",
            f"**原：** {old}",
            "",
            f"**新：** {new}",
            "",
        ]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


# ---------------------------------------------------------------- check

def numeric_key(s):
    """数字字符串按数值排序，避免 '10' 排在 '9' 前面"""
    try:
        return (0, float(s), s)
    except (TypeError, ValueError):
        return (1, 0.0, s)


def punct_count(t):
    return {
        "period": t.count("。"),
        "comma": t.count("，"),
        "semicolon": t.count("；"),
    }


def strip_cites(t):
    """算数字前后先去掉引用标记，避免 [1] 里的 1 被当成正文数字重复报警"""
    return CITE_RE.sub("", t)


def cite_numbers(t):
    """抽出引用标记里的编号，只比编号不比方括号样式。

    「[1,2]」写成「[1，2]」不算丢失；「[2-4]」展开成 2、3、4，
    这样改写成「[2,3,4]」也不算丢失。
    """
    out = []
    for m in CITE_RE.finditer(t):
        token = re.sub(r"^[\[［【]|[\]］】]$", "", m.group(0))
        prev, pending = None, False
        for num, sep in re.findall(r"(\d+)\s*([-,，–—~]?)", token):
            n = int(num)
            if pending and prev is not None and prev < n <= prev + 50:
                out += list(range(prev + 1, n))
            out.append(n)
            prev = n
            pending = bool(sep) and sep in "-–—~"
    return out


def echo_hits(t):
    """新文本里是否夹带了指令回显（模型没有只返回改写后的段落）"""
    return [m for m in ECHO_MARKERS if m in t]


def cmd_check(args):
    verify_blocks_source(args.blocks, args.docx)
    verify_edits_source(args.edits, args.docx)
    doc = open_docx(args.docx)
    blocks = scan_blocks(doc)
    block_map = {b["id"]: b for b in blocks}
    edits = load_edits(args.edits)
    paras = doc.paragraphs

    out_of_range, blocked = check_targets(paras, edits, block_map, args.allow_non_para)

    old_all, new_all = [], []
    per_para = []
    # write 会跳过含超链接的段落，check 必须同样跳过，否则自检数字与成品不一致
    skipped_link = []
    echo_ids = []
    for pid in sorted(edits):
        if not (0 <= pid < len(paras)):
            continue
        if any(pid == b[0] for b in blocked):
            continue
        p = paras[pid]
        if "<w:hyperlink" in p._element.xml and not args.force:
            skipped_link.append(pid)
            continue
        old = (p.text or "").strip()
        new = edits[pid].strip()
        hits = echo_hits(new)
        if hits:
            echo_ids.append((pid, hits))
        old_all.append(old)
        new_all.append(new)
        po, pn = punct_count(old), punct_count(new)
        old_nums = Counter(NUM_RE.findall(strip_cites(old)))
        new_nums = Counter(NUM_RE.findall(strip_cites(new)))
        old_cites, new_cites = Counter(cite_numbers(old)), Counter(cite_numbers(new))
        old_values = Counter(re.sub(r"\s+", "", n) for n in VALUE_RE.findall(strip_cites(old)))
        new_values = Counter(re.sub(r"\s+", "", n) for n in VALUE_RE.findall(strip_cites(new)))
        per_para.append(
            {
                "id": pid,
                "chars_old": len(old),
                "chars_new": len(new),
                "chars_pct": round((len(new) - len(old)) / max(len(old), 1) * 100, 1),
                "period_old": po["period"],
                "period_new": pn["period"],
                "numbers_lost": sorted(set(old_nums) - set(new_nums), key=numeric_key),
                "numbers_reduced": sorted(
                    f"{k}×{v}→×{new_nums.get(k, 0)}"
                    for k, v in old_nums.items()
                    if new_nums.get(k, 0) < v and k in new_nums
                ),
                "cites_lost": sorted((old_cites - new_cites).elements()),
                "cites_added": sorted((new_cites - old_cites).elements()),
                "values_changed": old_values != new_values,
                "banned": [w for w in BANNED_SUBJECTIVE if w in new],
                "echo": hits,
            }
        )

    old_text, new_text = "".join(old_all), "".join(new_all)
    po, pn = punct_count(old_text), punct_count(new_text)

    all_nums_lost, all_nums_reduced, all_cites_lost, banned_hits = [], [], [], []
    for r in per_para:
        all_nums_lost += r["numbers_lost"]
        all_nums_reduced += r["numbers_reduced"]
        all_cites_lost += r["cites_lost"]
        banned_hits += r["banned"]

    char_pct = round((len(new_text) - len(old_text)) / max(len(old_text), 1) * 100, 1)
    period_pct = (
        round((pn["period"] - po["period"]) / max(po["period"], 1) * 100, 1)
        if po["period"]
        else 0.0
    )

    issues, warnings = [], []
    for r in per_para:
        if r["cites_added"]:
            issues.append(f"段落 {r['id']} 新增引用标记：{r['cites_added']}")
        if r["values_changed"]:
            issues.append(f"段落 {r['id']} 数值、符号、百分号或出现次数改变，需人工核对")
    if blocked:
        detail = "；".join(f"{i}({KIND_LABEL.get(k, k)})" for i, k, _ in blocked)
        issues.append(f"{len(blocked)} 个 id 不是正文段落，未计入自检：{detail}")
    if out_of_range:
        warnings.append(f"{len(out_of_range)} 个段落 id 越界，已忽略：{out_of_range[:10]}")
    if skipped_link:
        warnings.append(
            f"{len(skipped_link)} 段含超链接，write 会跳过它们，本次未计入指标：{skipped_link[:10]}"
        )
    if echo_ids:
        detail = "；".join(f"{i}({','.join(h)})" for i, h in echo_ids)
        issues.append(f"{len(echo_ids)} 段把改写指令回显进了正文，必须重新生成：{detail}")

    if per_para:
        if abs(char_pct) > 10:
            warnings.append(f"字数变化 {char_pct}%，超出 ±10% 口径")
        if all_cites_lost:
            marks = "、".join(f"[{n}]" for n in all_cites_lost[:8])
            issues.append(f"引用标记丢失 {len(all_cites_lost)} 处：{marks}")
        if all_nums_lost:
            warnings.append(f"数字未原样出现 {len(all_nums_lost)} 处：{all_nums_lost[:8]}")
        if all_nums_reduced:
            warnings.append(f"数字出现次数减少 {len(all_nums_reduced)} 处：{all_nums_reduced[:6]}")
        if banned_hits:
            issues.append(f"残留主观表述：{sorted(set(banned_hits))}")
        if po["period"] and period_pct > -20:
            warnings.append(f"句号仅减少 {abs(period_pct)}%，提示词要求「减少句号使用」")
        if pn["comma"] <= po["comma"]:
            warnings.append("逗号数未增加，未体现「多用逗号衔接」")
    else:
        warnings.append("没有可自检的段落：edits 为空，或所有 id 都被拦截")

    verdict = "fail" if issues else ("warn" if warnings else "pass")

    out = {
        "source": os.path.abspath(args.docx),
        "checked_at": datetime.now().isoformat(timespec="seconds"),
        "paras_edited": len(per_para),
        "chars": {"old": len(old_text), "new": len(new_text), "pct": char_pct},
        "punct": {
            "period": {"old": po["period"], "new": pn["period"], "pct": period_pct},
            "comma": {"old": po["comma"], "new": pn["comma"]},
            "semicolon": {"old": po["semicolon"], "new": pn["semicolon"]},
        },
        "numbers_lost": all_nums_lost,
        "numbers_reduced": all_nums_reduced,
        "cites_lost": all_cites_lost,
        "banned_hits": sorted(set(banned_hits)),
        "ai_cliche_left": [w for w in AI_CLICHE if w in new_text],
        "echo_ids": [i for i, _ in echo_ids],
        "blocked_ids": [b[0] for b in blocked],
        "out_of_range_ids": out_of_range,
        "skipped_link_ids": skipped_link,
        "issues": issues,
        "warnings": warnings,
        "verdict": verdict,
        "per_para": per_para,
    }

    if args.output:
        write_json(args.output, out)
    print(f"[check] {verdict.upper()}  改动 {len(per_para)} 段")
    print(f"  字数  {len(old_text)} -> {len(new_text)}  ({char_pct:+.1f}%)")
    print(f"  句号  {po['period']} -> {pn['period']}  ({period_pct:+.1f}%)   逗号 {po['comma']} -> {pn['comma']}")
    for x in issues:
        print(f"  [X] {x}")
    for x in warnings:
        print(f"  [!] {x}")
    if args.output:
        print(f"  报告 -> {args.output}")
    if verdict == "fail":
        raise SystemExit(1)


# ---------------------------------------------------------------- apply / status

def read_text(path):
    """读模型回复：优先 UTF-8，Windows 上常见 GB18030 也一并兼容"""
    if not os.path.exists(path):
        sys.exit(f"[错误] 文件不存在：{path}")
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    sys.exit(f"[错误] 读不出 {os.path.basename(path)} 的编码，请另存为 UTF-8 再试")


def split_reply(text, strip_number=False):
    """把模型回复切成段落块。

    规则：按空行分块；块内硬换行直接合并（中文正文不应带硬换行）；
    先剔除整段包裹的 ``` 代码围栏。切出来的块数必须与段落数一致。
    """
    t = text.replace("\r\n", "\n").replace("\r", "\n")
    t = re.sub(r"(?m)^\s*```[^\n]*$", "", t)
    chunks = [c.strip() for c in re.split(r"\n\s*\n", t)]
    chunks = [c for c in chunks if c]
    if strip_number:
        chunks = [
            re.sub(r"^\s*(?:\d+\s*\.(?!\d)|\d+\s*[、)]|[（(]\s*\d+\s*[）)])\s*", "", c) for c in chunks
        ]
    return [re.sub(r"\s*\n\s*", "", c) for c in chunks if c.strip()]


def parse_ids_arg(raw):
    try:
        return [int(x) for x in raw.replace("，", ",").split(",") if x.strip()]
    except ValueError:
        sys.exit(f"[错误] --ids 只接受逗号分隔的整数，收到：{raw}")


def cmd_apply(args):
    data = load_json(args.blocks)
    verify_blocks_source(args.blocks, data.get("source"))
    batches = data["batches"]
    block_map = {b["id"]: b for b in data["blocks"]}

    if args.batch is not None:
        hit = [b for b in batches if b["batch_id"] == args.batch]
        if not hit:
            sys.exit(f"[错误] 没有第 {args.batch} 批，共 {len(batches)} 批")
        ids = list(hit[0]["para_ids"])
        label = f"批次 {args.batch}（{hit[0]['section']}）"
    elif args.ids:
        ids = parse_ids_arg(args.ids)
        label = "自定义段落"
    else:
        sys.exit("[错误] 需要 --batch 或 --ids 之一")

    if not ids:
        sys.exit("[错误] 这一批没有可改写的正文段落")
    if len(ids) != len(set(ids)):
        sys.exit("[错误] 段落 id 重复，拒绝处理")

    bad = [i for i in ids if i not in block_map]
    if bad:
        sys.exit(f"[错误] blocks 里不存在这些段落 id：{bad}")
    blocked = [i for i in ids if block_map[i]["kind"] not in EDITABLE_KINDS]
    if blocked:
        detail = "\n  ".join(
            f"{i} -> {KIND_LABEL.get(block_map[i]['kind'], block_map[i]['kind'])}：{block_map[i]['text'][:30]}"
            for i in blocked
        )
        sys.exit(f"[错误] 这些 id 不是正文段落，不能参与改写：\n  {detail}")

    reply = read_text(args.input) if args.input else sys.stdin.read()
    chunks = split_reply(reply, args.strip_number)

    if len(chunks) != len(ids):
        got = "\n  ".join(f"[{n + 1}] {c[:36]}…" for n, c in enumerate(chunks[:8])) or "  （空回复）"
        sys.exit(
            f"[错误] 段数对不上，拒绝写入。\n"
            f"  {label}：需要 {len(ids)} 段，回复里切出 {len(chunks)} 块。\n"
            f"  期望的段落 id：{','.join(str(i) for i in ids)}\n"
            f"  切出来的块（前 8 块）：\n  {got}\n"
            "  常见原因：模型把两段并成一段、多返回了说明文字、或段落之间漏了空行。\n"
            "  处理：重新生成该批；确需手工拼接时，先按空行把段落分开再重跑 apply。"
        )

    edits = {}
    if os.path.exists(args.edits):
        existing = load_json(args.edits)
        if not isinstance(existing, dict) or existing.get("source_sha256") != data["source_sha256"]:
            sys.exit("[错误] 已有 edits 来源不明或属于另一原稿，请使用新的工作目录")
        edits = {str(k): v for k, v in load_edits(args.edits).items()}
    before = len(edits)
    for pid, chunk in zip(ids, chunks):
        edits[str(pid)] = chunk

    drift = []
    for pid, chunk in zip(ids, chunks):
        old_len = block_map[pid]["chars"] or 1
        pct = (len(chunk) - old_len) / old_len * 100
        if abs(pct) > 15:
            drift.append(f"{pid}({pct:+.0f}%)")

    if args.dry_run:
        print(f"[apply] 试运行 {label}：{len(ids)} 段全部对齐，未写文件")
        for pid, chunk in zip(ids, chunks):
            print(f"  {pid}: {chunk[:48]}")
        return

    write_json(args.edits, {"schema_version": 1, "source_sha256": data["source_sha256"], "edits": edits})
    print(f"[apply] {label}：写入 {len(ids)} 段 -> {args.edits}（累计 {before} -> {len(edits)} 段）")
    if drift:
        print(
            f"  [提示] {len(drift)} 段字数偏离原文超过 ±15%：{drift[:8]}"
            "（改写到位的话应该接近原文长度）"
        )

    if args.progress:
        prog = {"total_batches": len(batches), "done": [], "failed": [], "updated_at": ""}
        if os.path.exists(args.progress):
            try:
                prog.update(load_json(args.progress))
            except SystemExit:
                print(f"  [提示] {args.progress} 读不出来，已重建进度文件", file=sys.stderr)
        prog["total_batches"] = len(batches)
        done = set(prog.get("done") or [])
        if args.batch is not None:
            done.add(args.batch)
            prog["failed"] = [x for x in (prog.get("failed") or []) if x != args.batch]
        prog["done"] = sorted(done)
        prog["updated_at"] = datetime.now().isoformat(timespec="seconds")
        write_json(args.progress, prog)
        print(f"  进度 {len(prog['done'])}/{len(batches)} 批 -> {args.progress}")


def cmd_status(args):
    data = load_json(args.blocks)
    edits = {}
    if os.path.exists(args.edits):
        edits = load_edits(args.edits)
    batches = data["batches"]
    done_batches = 0
    print(f"[status] {os.path.basename(data.get('source') or args.blocks)}  共 {len(batches)} 批")
    for b in batches:
        ids = b["para_ids"]
        got = [i for i in ids if i in edits]
        if len(got) == len(ids):
            state, done_batches = "完成", done_batches + 1
        elif got:
            state = "部分"
        else:
            state = "未开始"
        print(f"  批 {b['batch_id']:>2}  {state:<4} {len(got)}/{len(ids)} 段  {b['section'][:26]}")
    print(f"  已写入段落 {len(edits)} 段；完成 {done_batches}/{len(batches)} 批")
    if done_batches < len(batches):
        nxt = next(b["batch_id"] for b in batches
                   if len([i for i in b["para_ids"] if i in edits]) < len(b["para_ids"]))
        print(f"  下一批：--batch {nxt}")


# ---------------------------------------------------------------- utils

def open_docx(path):
    """打开 docx，把常见的打不开情况翻译成人话"""
    if not os.path.exists(path):
        sys.exit(f"[错误] 文件不存在：{path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".doc":
        sys.exit(
            f"[错误] {os.path.basename(path)} 是旧版 .doc 格式，python-docx 读不了。\n"
            "  请用 Word 打开后「另存为」→ Word 文档(*.docx)，再重试。"
        )
    try:
        return Document(path)
    except Exception as e:
        sys.exit(
            f"[错误] 打不开这个 docx：{path}\n"
            f"  原因：{type(e).__name__}: {e}\n"
            "  常见情况：文件损坏、设有打开密码、或者其实是改了扩展名的 .doc / PDF。\n"
            "  排查办法：用 Word 打开它，「另存为」一份新的 .docx 再试。"
        )


def load_json(path):
    if not os.path.exists(path):
        sys.exit(f"[错误] 文件不存在：{path}")
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        sys.exit(
            f"[错误] {os.path.basename(path)} 不是合法 JSON："
            f"第 {e.lineno} 行第 {e.colno} 列附近，{e.msg}"
        )
    except UnicodeDecodeError:
        sys.exit(f"[错误] {os.path.basename(path)} 不是 UTF-8 编码，请用 UTF-8 保存后再试")


def write_json(path, obj):
    folder = os.path.dirname(os.path.abspath(path)) or "."
    os.makedirs(folder, exist_ok=True)
    tmp = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=folder,
                                         suffix=".tmp", delete=False) as f:
            tmp = f.name
            json.dump(obj, f, ensure_ascii=False, indent=2)
        os.replace(tmp, path)
    finally:
        if tmp and os.path.exists(tmp):
            os.unlink(tmp)


def file_digest(path):
    digest = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError:
        sys.exit(f"[错误] 无法读取来源文件：{path}")
    return digest.hexdigest()


def verify_blocks_source(path, source):
    if not path:
        print("[提示] 未提供 --blocks；仅在新版 edits 携带指纹时自动验证来源", file=sys.stderr)
        return
    data = load_json(path)
    expected = data.get("source_sha256")
    if not expected:
        sys.exit("[错误] blocks 缺少来源指纹，请重新 split")
    if not source or not os.path.isfile(source) or file_digest(source) != expected:
        sys.exit("[错误] 原稿与 split 时的来源不一致，请重新 split 并重新生成改写")


def verify_edits_source(path, source):
    data = load_json(path)
    if isinstance(data, dict) and "schema_version" in data:
        if data.get("source_sha256") != file_digest(source):
            sys.exit("[错误] edits 的来源指纹与原稿不一致，拒绝套用其他文档的改写")
    else:
        print("[提示] 旧格式 edits 没有来源指纹，请配合 --blocks 并人工核对 id", file=sys.stderr)


def guard_output_paths(args):
    inputs = [getattr(args, k, None) for k in ("docx", "blocks", "input", "prompt")]
    if args.cmd in ("apply", "batch", "status"):
        data = load_json(args.blocks)
        if isinstance(data, dict) and isinstance(data.get("source"), str):
            inputs.append(data["source"])
    outputs = [getattr(args, k, None) for k in ("output", "diff_md", "progress")]
    if args.cmd in ("write", "check", "status"):
        inputs.append(args.edits)
    if args.cmd == "apply":
        outputs.append(args.edits)
    inputs = [p for p in inputs if p]
    outputs = [p for p in outputs if p]
    for i, out in enumerate(outputs):
        if any(same_file(out, p) for p in inputs + outputs[:i]):
            sys.exit("[错误] 输出路径与输入或其他输出重合，拒绝覆盖")


def main():
    ap = argparse.ArgumentParser(
        description="论文表达整改：split / batch / apply / status / write / check"
    )
    sub = ap.add_subparsers(dest="cmd", required=True)

    s1 = sub.add_parser("split", help="docx -> 段落清单 + 章节树 + 分批方案")
    s1.add_argument("docx")
    s1.add_argument("-o", "--output", default="work/blocks.json")
    s1.add_argument("--batch-chars", type=int, default=1200)
    s1.set_defaults(func=cmd_split)

    s2 = sub.add_parser("batch", help="取一批段落，拼成可直接发送的改写请求")
    s2.add_argument("blocks")
    s2.add_argument("--id", type=int, default=None, help="批次号（从 split 结果里取）")
    s2.add_argument("--ids", default=None, help="直接指定段落 id，逗号分隔，如 8,9,10")
    s2.add_argument("--section", default=None, help="按章节标题模糊匹配")
    s2.add_argument("--prompt", default=None, help="改写指令文件，如 references/prompt.md")
    s2.add_argument("-o", "--output", default=None, help="输出到文件（纯请求文本）")
    s2.set_defaults(func=cmd_batch)

    s3 = sub.add_parser("write", help="改写结果 -> 新 docx")
    s3.add_argument("docx")
    s3.add_argument("edits")
    s3.add_argument("--blocks", default=None, help="split 清单，用于验证原稿来源指纹")
    s3.add_argument("-o", "--output", required=True)
    s3.add_argument("--diff-md", default=None, help="可选：显式导出逐段对照表（默认不生成）")
    s3.add_argument("--force", action="store_true", help="强制改写含超链接的段落（会使链接文字重复，慎用）")
    s3.add_argument("--allow-non-para", action="store_true", help="允许改写标题/图注等非正文段落（默认拒绝）")
    s3.set_defaults(func=cmd_write)

    s4 = sub.add_parser("check", help="改写前后硬指标自检")
    s4.add_argument("docx")
    s4.add_argument("edits")
    s4.add_argument("--blocks", default=None, help="split 清单，用于验证原稿来源指纹")
    s4.add_argument("-o", "--output", default=None, help="可选：显式导出自检 JSON（默认仅内部检查，不生成文件）")
    s4.add_argument("--allow-non-para", action="store_true", help="允许把非正文段落计入自检（默认忽略）")
    s4.add_argument("--force", action="store_true", help="把含超链接的段落也计入自检（与 write --force 配套）")
    s4.set_defaults(func=cmd_check)

    s5 = sub.add_parser("apply", help="模型回复 -> edits.json（按段落数严格对齐）")
    s5.add_argument("blocks")
    s5.add_argument("--batch", type=int, default=None, help="批次号")
    s5.add_argument("--ids", default=None, help="直接指定段落 id，逗号分隔")
    s5.add_argument("--input", default=None, help="模型回复的文本文件；省略则从 stdin 读")
    s5.add_argument("-o", "--edits", default="work/edits.json", help="累积写入的 edits.json")
    s5.add_argument("--progress", default=None, help="回填进度，如 work/progress.json")
    s5.add_argument("--strip-number", action="store_true", help="去掉每块开头的「1.」「（1）」编号")
    s5.add_argument("--dry-run", action="store_true", help="只校验对齐，不写文件")
    s5.set_defaults(func=cmd_apply)

    s6 = sub.add_parser("status", help="查看各批完成情况，便于中断后续跑")
    s6.add_argument("blocks")
    s6.add_argument("-e", "--edits", default="work/edits.json")
    s6.set_defaults(func=cmd_status)

    args = ap.parse_args()
    guard_output_paths(args)
    args.func(args)


if __name__ == "__main__":
    main()
