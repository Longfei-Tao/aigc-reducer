#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""回归测试：守住「不许把稿子改坏」这条线。

跑法（在仓库根目录，不需要 pytest）：

    python -m unittest discover -s tests -v

覆盖范围：
    1. split 的段落分类（含中文编号标题、年份开头的正文）
    2. batch 添加执行约束，并完整保留 references/prompt.md 的风格指令与原文
    3. write 的四类拦截与排版保留
    4. apply 的段落对齐，段数不符必须拒绝
    5. check 的引用识别、口径一致性、指令回显判定
"""

import json
import os
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "aigc-reducer" / "scripts"
TOOL = SCRIPTS / "docx_tool.py"
PROMPT = ROOT / "skills" / "aigc-reducer" / "references" / "prompt.md"

sys.path.insert(0, str(SCRIPTS))

import docx_tool  # noqa: E402

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

# 旧版 Windows 终端默认 GBK，脚本输出是 UTF-8，这里统一按 UTF-8 解，避免测试假失败
ENV = dict(os.environ, PYTHONIOENCODING="utf-8")


def run(*args, cwd=None, stdin=None):
    """跑一次命令行，返回 (退出码, stdout, stderr)"""
    p = subprocess.run(
        [sys.executable, str(TOOL), *[str(a) for a in args]],
        cwd=str(cwd) if cwd else None,
        input=stdin,
        capture_output=True,
        env=ENV,
    )
    return (
        p.returncode,
        p.stdout.decode("utf-8", "replace"),
        p.stderr.decode("utf-8", "replace"),
    )


def make_doc(path, texts, hyperlink_at=None, heading_at=()):
    """造一份小 docx：texts 顺序写入，可指定哪些下标是标题、哪一段挂超链接"""
    doc = Document()
    for i, t in enumerate(texts):
        p = doc.add_paragraph()
        if hyperlink_at == i:
            p.add_run(t[:2])
            hl = OxmlElement("w:hyperlink")
            hl.set(qn("r:id"), "rId999")
            r = OxmlElement("w:r")
            node = OxmlElement("w:t")
            node.text = t[2:6]
            r.append(node)
            hl.append(r)
            p._p.append(hl)
            p.add_run(t[6:])
        else:
            p.add_run(t)
        if i in heading_at:
            p.style = doc.styles["Heading 1"]
    doc.save(str(path))


class Base(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.tmp.name)
        cls.sample = cls.dir / "sample.docx"
        subprocess.run(
            [sys.executable, str(SCRIPTS / "make_sample.py"), "-o", str(cls.sample)],
            check=True,
            capture_output=True,
            env=ENV,
        )
        cls.blocks = cls.dir / "work" / "blocks.json"
        code, out, err = run("split", cls.sample, "-o", cls.blocks)
        assert code == 0, err

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def load_blocks(self, path=None):
        with open(path or self.blocks, encoding="utf-8") as f:
            return json.load(f)


class TestSplit(Base):
    def test_summary_numbers(self):
        data = self.load_blocks()
        s = data["summary"]
        self.assertEqual(s["para_total"], 25)
        self.assertEqual(s["eligible_paras"], 10)
        self.assertEqual(s["tables"], 1)
        self.assertGreaterEqual(s["batches"], 5)

    def test_batches_never_cross_heading(self):
        data = self.load_blocks()
        kind = {b["id"]: b["kind"] for b in data["blocks"]}
        for batch in data["batches"]:
            for pid in batch["para_ids"]:
                self.assertEqual(kind[pid], "para", f"批 {batch['batch_id']} 混入了非正文段落")

    def test_reference_section_detected(self):
        data = self.load_blocks()
        refs = [b for b in data["blocks"] if b["kind"] == "ref"]
        self.assertEqual(len(refs), 4)  # 3 条文献 + 致谢正文

    def test_chinese_heading_variants(self):
        doc = self.dir / "headings.docx"
        texts = ["一、研究背景", "（一）研究背景", "第一章 绪论", "2.3.1 统计方法"]
        make_doc(doc, texts)
        out = self.dir / "headings.json"
        code, _, err = run("split", doc, "-o", out)
        self.assertEqual(code, 0, err)
        kinds = [b["kind"] for b in self.load_blocks(out)["blocks"]]
        self.assertEqual(kinds, ["heading"] * 4, f"中文编号标题未识别：{kinds}")

    def test_year_start_paragraph_is_body(self):
        doc = self.dir / "year.docx"
        make_doc(doc, ["2023 年共收集 186 份问卷，有效 180 份"])
        out = self.dir / "year.json"
        run("split", doc, "-o", out)
        self.assertEqual(self.load_blocks(out)["blocks"][0]["kind"], "para")


class TestBatch(Base):
    def test_payload_keeps_verbatim_prompt_and_source_after_constraints(self):
        out = self.dir / "batch1.txt"
        code, _, err = run("batch", self.blocks, "--id", 1, "--prompt", PROMPT, "-o", out)
        self.assertEqual(code, 0, err)
        raw = PROMPT.read_text(encoding="utf-8")
        body = raw[re.search(r"(?m)^---\s*$", raw).end():].lstrip("\n").rstrip() + "\n"
        payload = out.read_text(encoding="utf-8")
        self.assertIn("【改写指令】\n", payload)
        constraints, request = payload.split("【改写指令】\n", 1)
        self.assertIn("不调用联网搜索", constraints)
        self.assertIn("不能切换或确认", constraints)
        self.assertIn("只返回改写后的段落", constraints)
        data = self.load_blocks()
        ids = data["batches"][0]["para_ids"]
        source = "\n\n".join(b["text"].strip() for i in ids for b in data["blocks"] if b["id"] == i)
        self.assertEqual(request, body + source + "\n", "风格指令或原文被改动")

    def test_raw_batch_without_prompt_stays_source_only(self):
        code, payload, err = run("batch", self.blocks, "--id", 1)
        self.assertEqual(code, 0, err)
        data = self.load_blocks()
        ids = data["batches"][0]["para_ids"]
        source = "\n\n".join(b["text"].strip() for i in ids for b in data["blocks"] if b["id"] == i)
        self.assertEqual(payload.replace("\r\n", "\n"), source + "\n")

    def test_custom_prompt_also_gets_constraints(self):
        custom = self.dir / "custom_prompt.md"
        custom.write_text("说明\n---\n自定义风格，保留术语。\n修改段落:\n", encoding="utf-8")
        code, payload, err = run("batch", self.blocks, "--id", 1, "--prompt", custom)
        self.assertEqual(code, 0, err)
        payload = payload.replace("\r\n", "\n")
        self.assertIn("【改写指令】\n", payload)
        constraints, request = payload.split("【改写指令】\n", 1)
        self.assertIn("不调用联网搜索", constraints)
        self.assertTrue(request.startswith("自定义风格，保留术语。\n修改段落:\n"))

    def test_rejects_non_para_ids(self):
        data = self.load_blocks()
        heading = next(b["id"] for b in data["blocks"] if b["kind"] == "heading")
        code, out, err = run("batch", self.blocks, "--ids", heading)
        self.assertEqual(code, 1)
        self.assertIn("不是正文段落", err)
        self.assertEqual(out, "")


class TestWrite(Base):
    def test_refuses_to_overwrite_source(self):
        edits = self.dir / "e_same.json"
        edits.write_text('{"1": "x"}', encoding="utf-8")
        code, _, err = run("write", self.sample, edits, "-o", self.sample)
        self.assertEqual(code, 1)
        self.assertIn("拒绝覆盖", err)

    def test_blocks_non_para_target(self):
        edits = self.dir / "e_head.json"
        edits.write_text('{"3": "被覆盖的标题"}', encoding="utf-8")
        out = self.dir / "should_not_exist.docx"
        code, _, err = run("write", self.sample, edits, "-o", out)
        self.assertEqual(code, 1)
        self.assertIn("不是正文段落", err)
        self.assertFalse(out.exists(), "拦截后不应生成文件")

    def test_blocks_prompt_echo(self):
        edits = self.dir / "e_echo.json"
        for index, marker in enumerate(("修改段落:", "【执行约束】", "【改写指令】")):
            with self.subTest(marker=marker):
                output = self.dir / f"echo_{index}.docx"
                edits.write_text(json.dumps({"1": marker + "本研究旨在探讨术后早期营养干预。"}, ensure_ascii=False), encoding="utf-8")
                code, _, err = run("write", self.sample, edits, "-o", output)
                self.assertEqual(code, 1)
                self.assertIn("回显", err)
                self.assertFalse(output.exists())

    def test_preserves_styles_table_and_untouched_paras(self):
        edits = self.dir / "e_ok.json"
        edits.write_text(
            json.dumps({str(i): f"改写后的第 {i} 段。" for i in (1, 2, 4, 5)}, ensure_ascii=False),
            encoding="utf-8",
        )
        out = self.dir / "ok.docx"
        code, _, err = run("write", self.sample, edits, "-o", out)
        self.assertEqual(code, 0, err)
        a, b = Document(str(self.sample)), Document(str(out))
        self.assertEqual(len(a.paragraphs), len(b.paragraphs))
        self.assertEqual(
            [p.style.name for p in a.paragraphs], [p.style.name for p in b.paragraphs]
        )
        self.assertEqual(
            [[c.text for c in r.cells] for r in a.tables[0].rows],
            [[c.text for c in r.cells] for r in b.tables[0].rows],
        )
        kept = [0, 3, 6, 7, 9, 11, 12, 14, 16, 19, 20, 21, 22, 23, 24]
        for i in kept:
            self.assertEqual(a.paragraphs[i].text, b.paragraphs[i].text, f"第 {i} 段不该被改动")

    def test_hyperlink_paragraph_skipped(self):
        doc = self.dir / "link.docx"
        make_doc(doc, ["参见国家统计局网站的公开数据，口径与其保持一致。"], hyperlink_at=0)
        blocks = self.dir / "link.json"
        run("split", doc, "-o", blocks)
        edits = self.dir / "e_link.json"
        edits.write_text('{"0": "参见国家统计局的公开数据，口径保持一致。"}', encoding="utf-8")
        code, out, err = run("write", doc, edits, "-o", self.dir / "link_out.docx")
        self.assertEqual(code, 1, "唯一一段被跳过时不应生成文件")
        self.assertIn("超链接", err)


class TestApply(Base):
    def test_aligns_chunks_to_ids(self):
        out = self.dir / "e_apply.json"
        reply = "第一段改写文字。\n\n第二段改写文字。\n"
        code, _, err = run(
            "apply", self.blocks, "--batch", 1, "-o", out, stdin=reply.encode("utf-8")
        )
        self.assertEqual(code, 0, err)
        data = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(data["source_sha256"], self.load_blocks()["source_sha256"])
        self.assertEqual(sorted(data["edits"]), ["1", "2"])
        self.assertEqual(data["edits"]["1"], "第一段改写文字。")

    def test_rejects_count_mismatch(self):
        out = self.dir / "e_apply_bad.json"
        code, _, err = run(
            "apply", self.blocks, "--batch", 1, "-o", out, stdin="只有一段。\n".encode("utf-8")
        )
        self.assertEqual(code, 1)
        self.assertIn("段数对不上", err)
        self.assertFalse(out.exists(), "段数不符时不应写入 edits.json")

    def test_accumulates_and_strips_fence(self):
        out = self.dir / "e_accum.json"
        run("apply", self.blocks, "--batch", 1, "-o", out, stdin="甲。\n\n乙。\n".encode("utf-8"))
        run("apply", self.blocks, "--batch", 3, "-o", out, stdin="```\n丙。\n```\n".encode("utf-8"))
        data = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(sorted(data["edits"], key=int), ["1", "2", "8"])
        self.assertEqual(data["edits"]["8"], "丙。")

    def test_progress_written(self):
        out = self.dir / "e_prog.json"
        prog = self.dir / "progress.json"
        run(
            "apply", self.blocks, "--batch", 1, "-o", out, "--progress", prog,
            stdin="甲。\n\n乙。\n".encode("utf-8"),
        )
        data = json.loads(prog.read_text(encoding="utf-8"))
        self.assertEqual(data["done"], [1])
        self.assertEqual(data["total_batches"], self.load_blocks()["summary"]["batches"])


class TestCheck(Base):
    def test_parenthetical_number_is_not_citation(self):
        self.assertEqual(docx_tool.cite_numbers("国内学者（2021）对此有过讨论"), [])
        self.assertEqual(docx_tool.cite_numbers("结果见表（1）"), [])
        self.assertEqual(docx_tool.cite_numbers("文献[2-4]"), [2, 3, 4])

    def test_citation_loss_is_fail(self):
        edits = self.dir / "e_cite.json"
        edits.write_text(
            '{"2": "从结果看，干预组的平均住院天数为 7.2 天，对照组为 10.8 天，'
            '两组差异有统计学意义（P<0.05），说明干预组恢复得更快。"}',
            encoding="utf-8",
        )
        out = self.dir / "c_cite.json"
        code, _, err = run("check", self.sample, edits, "-o", out)
        self.assertEqual(code, 1, err)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "fail")
        self.assertEqual(report["cites_lost"], [2])

    def test_echo_is_fail(self):
        edits = self.dir / "e_echo2.json"
        edits.write_text(
            '{"4": "指令:1. 在不大幅改动原文总字数的情况下。修改段落:综上所论。"}',
            encoding="utf-8",
        )
        out = self.dir / "c_echo.json"
        run("check", self.sample, edits, "-o", out)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "fail")
        self.assertIn(4, report["echo_ids"])

    def test_hyperlink_scope_matches_write(self):
        doc = self.dir / "link2.docx"
        make_doc(doc, ["参见国家统计局网站的公开数据，口径与其保持一致。"], hyperlink_at=0)
        blocks = self.dir / "link2.json"
        run("split", doc, "-o", blocks)
        edits = self.dir / "e_link2.json"
        edits.write_text('{"0": "参见国家统计局的公开数据，口径保持一致。"}', encoding="utf-8")
        out = self.dir / "c_link2.json"
        run("check", doc, edits, "-o", out)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["paras_edited"], 0, "超链接段不应计入自检")
        self.assertEqual(report["skipped_link_ids"], [0])

    def test_semicolon_increase_warns(self):
        doc = self.dir / "semi.docx"
        make_doc(doc, ["纳入标准为年龄十八岁以上；排除标准为合并严重器质性疾病；所有患者均签署知情同意书。"])
        edits = self.dir / "e_semi.json"
        edits.write_text(
            '{"0": "纳入标准是年龄十八岁以上；排除标准是合并严重器质性疾病；'
            '所有患者都签了知情同意书；分组按入院顺序确定。"}',
            encoding="utf-8",
        )
        out = self.dir / "c_semi.json"
        code, _, err = run("check", doc, edits, "-o", out)
        self.assertEqual(code, 0, err)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "warn")
        self.assertEqual(report["punct"]["semicolon"], {"old": 2, "new": 3, "pct": 50.0})
        self.assertTrue(any("分号" in w for w in report["warnings"]), report["warnings"])

    def test_connector_decrease_warns(self):
        doc = self.dir / "comma.docx"
        make_doc(doc, ["干预组每日补充两次，对照组按常规流程处理，两组的住院时间都记录下来。"])
        edits = self.dir / "e_comma.json"
        edits.write_text('{"0": "干预组每日补充两次，对照组按常规流程处理。"}', encoding="utf-8")
        out = self.dir / "c_comma.json"
        run("check", doc, edits, "-o", out)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "warn")
        self.assertTrue(any("逗号与句号合计" in w for w in report["warnings"]), report["warnings"])

    def test_splitting_sentences_is_not_penalized(self):
        # 拆短句会把逗号改成句号：合计不变时不应报错，否则合格的拆句会被误判
        doc = self.dir / "split_sent.docx"
        src = "干预组每天补充两次营养剂，对照组按常规流程处理，两组的住院时间都做了记录。"
        new = "干预组每天补充两次营养剂。对照组按常规流程处理。两组的住院时间都做了记录。"
        make_doc(doc, [src])
        edits = self.dir / "e_split.json"
        edits.write_text(json.dumps({"0": new}, ensure_ascii=False), encoding="utf-8")
        out = self.dir / "c_split.json"
        run("check", doc, edits, "-o", out)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "pass", report["warnings"])

    def test_sample_regression_passes(self):
        edits = self.dir / "e_pass.json"
        edits.write_text(
            json.dumps(
                {
                    "1": "本研究想看看术后早期营养干预对加速康复外科效果有没有影响，通过对 2023 年 6 月至 "
                    "2024 年 5 月期间收治的 186 例患者做回顾性分析，早期营养干预能够提升患者的康复质量，"
                    "也能缩短住院时间 [1]，这一点在数据里比较明显。",
                    "2": "从结果看，干预组的平均住院天数为 7.2 天，对照组为 10.8 天，"
                    "两组差异有统计学意义（P<0.05）[2]，说明干预组恢复得更快一些。",
                    "4": "加速康复外科（ERAS）这个提法出现以后，已经在多个外科领域用得比较多，"
                    "营养管理是 ERAS 里比较关键的一环，它的重要性这些年越来越明显，"
                    "因此找一套适合国人的营养干预方案就比较重要。",
                    "5": "目前国内的相关研究还比较少，多数研究样本量偏小，随访时间也不太长，"
                    "所以结论能推广的范围有限。",
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        out = self.dir / "c_pass.json"
        code, _, err = run("check", self.sample, edits, "-o", out)
        self.assertEqual(code, 0, err)
        report = json.loads(out.read_text(encoding="utf-8"))
        self.assertEqual(report["verdict"], "pass", report["issues"] + report["warnings"])
        self.assertEqual(report["paras_edited"], 4)


if __name__ == "__main__":
    unittest.main(verbosity=2)
